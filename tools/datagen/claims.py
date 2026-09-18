"""The claim table: every figure the workbook asserts, with the cell it sits in.

Before this module the mapping from a metric to a cell was implicit in `render_workbook`, which
walked `truth_metrics.json` and emitted cells inline. Nothing could be derived from that, because
there was no value to derive from: the coordinate existed only for as long as the write took.

Lifting it out is what makes the whole story possible. A mutation names a **claim**
(`guests_by_nationality`, `2026-01`, `DE`) and never a cell reference; the table resolves it to
`Nationality!B10`. Change the sheet layout and the spec keeps working, because the expectation is
built from the same table the renderer wrote from. A hand-written expectation citing `B10` would
quietly start asserting the wrong cell, which is the drift PRD-94's acceptance criterion names.

## What is in the table, and what deliberately is not

Only **verifiable claims**: the three occupancy columns across four periods, and the nationality
cross-tab. Not the `Total` row, which `tda.excel.read` treats as a `StatedTotal` and never compares
against truth. Not the `Rate & Revenue` sheet, which is out of scope by D-SCOPE-02. A cell that no
finding can ever be raised about has no business in a table whose purpose is deriving findings.

A country absent from a month has **no entry**, rather than an entry worth `None`. Truth has no key
for it either, so the set difference in `derive.py` sees nothing on either side and raises nothing.
That is the difference between a country that did not travel in February and a row the hotel
deleted: the second removes a cell that truth still has a key for, and becomes a V5.
"""

from __future__ import annotations

from dataclasses import dataclass, replace
from typing import TYPE_CHECKING, Final

from openpyxl.utils import get_column_letter

from datagen.aggregate import month_label
from datagen.spec import NATIONALITY_LABELS, QUARTER, WORKBOOK

if TYPE_CHECKING:
    from collections.abc import Iterable

# The three occupancy columns, in the order the sheet prints them. Column 1 is the period label.
OCCUPANCY_METRICS: Final[tuple[str, ...]] = (
    "room_nights_sold",
    "room_nights_available",
    "occupancy_pct",
)

NATIONALITY_METRIC: Final = "guests_by_nationality"
NATIONALITY_DIMENSION: Final = "nationality_iso2"


@dataclass(frozen=True, slots=True)
class ClaimCell:
    """One figure the hotel asserts, and where it sits.

    `label` and `label_cell` describe the *axis* cell the figure sits against — a country name in
    column A, or a month name. Several claims share one label cell, which is why a relabelling
    mutation rewrites the label on every claim carrying it rather than editing a coordinate.
    """

    metric: str
    period: str
    dimension: str | None
    value_key: str | None
    label: str
    label_cell: str
    sheet: str
    cell: str
    value: float

    @property
    def key(self) -> str:
        """The metric key this claim is compared against, in `MetricKey.rendered`'s format.

        The third statement of one format, after `aggregate._render_key` which writes the keys and
        `tda.contracts.MetricKey` which parses them. Third on purpose: this package may not import
        `tda.contracts`, and a key built here has to find the truth value the product will compare
        against. A divergence surfaces as "expected key absent from the verdict", which is loud.

        `test_every_baseline_claim_carries_a_truth_key_and_the_truth_value` in
        `tests/unit/test_datagen.py` pins all three together against every key in
        `truth_metrics.json`, from `tests/`, where importing both sides is allowed.
        """
        if self.dimension is None:
            return f"{self.metric}:{self.period}"
        return f"{self.metric}:{self.period}:{self.dimension}={self.value_key}"


def nationalities_in(truth: dict[str, float]) -> tuple[str, ...]:
    """The countries the sheet shows, in the order it shows them.

    Sorted by the printed label rather than the ISO code, because that is how a hotel orders a table
    a human reads. One consequence worth keeping: the workbook's row order then matches nothing in
    `truth_metrics.json`, so the reconciliation join has to be by key rather than by position.
    """
    present = {
        key.rsplit("=", 1)[1]
        for key in truth
        if key.startswith(f"{NATIONALITY_METRIC}:") and f":{NATIONALITY_DIMENSION}=" in key
    }
    return tuple(sorted(present, key=lambda iso2: NATIONALITY_LABELS[iso2]))


def baseline_table(months: tuple[str, ...], truth: dict[str, float]) -> tuple[ClaimCell, ...]:
    """Every claim the unmutated demo workbook makes, with its coordinates.

    The periods run months-then-quarter on the occupancy sheet and months-then-quarter across the
    nationality sheet's columns, matching what `render_workbook` prints. Both orders live here so
    the renderer and the derivation cannot disagree about them.
    """
    header = WORKBOOK.header_row
    cells: list[ClaimCell] = []

    for offset, period in (*enumerate(months), (len(months), QUARTER)):
        row = header + 1 + offset
        label = f"{QUARTER} total" if period == QUARTER else month_label(period)
        for index, metric in enumerate(OCCUPANCY_METRICS, start=2):
            cells.append(
                ClaimCell(
                    metric=metric,
                    period=period,
                    dimension=None,
                    value_key=None,
                    label=label,
                    label_cell=f"A{row}",
                    sheet=WORKBOOK.occupancy_sheet,
                    cell=f"{get_column_letter(index)}{row}",
                    value=truth[f"{metric}:{period}"],
                )
            )

    for offset, iso2 in enumerate(nationalities_in(truth)):
        row = header + 1 + offset
        label_cell = f"A{row}"
        for index, period in enumerate((*months, QUARTER), start=2):
            key = f"{NATIONALITY_METRIC}:{period}:{NATIONALITY_DIMENSION}={iso2}"
            # Absent means the country did not travel that month. No cell, rather than a cell worth
            # nothing: see the module docstring on why the distinction carries the whole of F4.
            if key not in truth:
                continue
            cells.append(
                ClaimCell(
                    metric=NATIONALITY_METRIC,
                    period=period,
                    dimension=NATIONALITY_DIMENSION,
                    value_key=iso2,
                    label=NATIONALITY_LABELS[iso2],
                    label_cell=label_cell,
                    sheet=WORKBOOK.nationality_sheet,
                    cell=f"{get_column_letter(index)}{row}",
                    value=truth[key],
                )
            )

    return tuple(cells)


def by_coordinate(table: Iterable[ClaimCell]) -> dict[tuple[str, str], ClaimCell]:
    """Indexed by `(sheet, cell)`, which is how the renderer looks a value up."""
    return {(cell.sheet, cell.cell): cell for cell in table}


def value_keys_in_order(table: Iterable[ClaimCell]) -> tuple[str, ...]:
    """The dimension values, in the row order the table already carries.

    The renderer takes its row order from here rather than recomputing it from `truth`, so a
    mutation that removes a country moves the `Total` row up by exactly one row without the
    renderer needing to know a mutation happened.
    """
    seen: dict[str, None] = {}
    for cell in table:
        if cell.value_key is not None:
            seen.setdefault(cell.value_key, None)
    return tuple(seen)


def by_key(table: Iterable[ClaimCell]) -> dict[str, ClaimCell]:
    """Indexed by metric key, which is how the derivation compares against truth."""
    return {cell.key: cell for cell in table}


def labels_by_value_key(table: Iterable[ClaimCell]) -> dict[str, str]:
    """The printed label for each dimension value, as the mutated table has it.

    The derivation needs this to say what an unresolvable label was, and the renderer needs it to
    print column A. Reading it off the table rather than off `NATIONALITY_LABELS` is what makes a
    relabelling mutation visible to both.
    """
    return {cell.value_key: cell.label for cell in table if cell.value_key is not None}


def relabelled(table: Iterable[ClaimCell], value_key: str, label: str) -> tuple[ClaimCell, ...]:
    """Every claim for `value_key`, carrying a new printed label. Coordinates unchanged."""
    return tuple(
        replace(cell, label=label) if cell.value_key == value_key else cell for cell in table
    )


def without(table: Iterable[ClaimCell], value_key: str) -> tuple[ClaimCell, ...]:
    """Every claim except those for `value_key`. The hotel omitted the row entirely."""
    return tuple(cell for cell in table if cell.value_key != value_key)
