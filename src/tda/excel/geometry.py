"""A1 ranges as a checked value object.

The mapping agent answers in A1 notation — `'B5:E25'`, `'A4'` — and every one of those strings is
**untrusted input**. That is not a statement about models in particular; a range typed into a
configuration file by a human would get the same treatment. What makes it worth a module is that the
failure is silent if it is not checked: `'B5:E26'` where `'B5:E25'` was meant reads the total row as
a country and produces a claim for a nationality called `Total`, which then fails to resolve, which
then surfaces three layers away as an unmappable country label. The cause and the symptom would be
nowhere near each other.

So a range is parsed once, into a rectangle whose bounds are known to be ordered, and every
downstream use takes the rectangle rather than the string.

**Why A1 strings rather than integers.** `AgentOutput` may not carry numeric fields
(`tools/guard/agent_schema_lint.py`), and a header row would otherwise be `header_row: int`. Working
in A1 notation is not a workaround for that rule — it is the rule pointing at the better design.
`'Nationality!D14'` is what `ExcelRef.citation` already renders, what a reviewer types into a
spreadsheet to go look, and what a cassette diff shows legibly. Row and column integers would have
had to be converted to it anyway, at every call site.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from typing import Final

from openpyxl.utils import get_column_letter, range_boundaries

# A single cell or a rectangular range, uppercase, no sheet qualifier and no `$`. Absolute markers
# are rejected rather than stripped: `$B$5` in a mapping means something produced a formula
# reference, and quietly normalising it would hide whatever that was.
_RANGE: Final = re.compile(r"^[A-Z]{1,3}[1-9][0-9]{0,6}(:[A-Z]{1,3}[1-9][0-9]{0,6})?$")


class RangeError(ValueError):
    """A range string is malformed, or does not fit the sheet it was mapped against."""


@dataclass(frozen=True, slots=True)
class Rect:
    """A rectangle of cells, 1-indexed and inclusive on all four sides.

    Bounds are ordered at construction, so nothing downstream has to ask whether `min_row` really is
    the smaller one.
    """

    min_col: int
    min_row: int
    max_col: int
    max_row: int

    def __post_init__(self) -> None:
        if self.min_col > self.max_col or self.min_row > self.max_row:
            raise RangeError(f"range bounds are inverted: {self!r}")
        if self.min_col < 1 or self.min_row < 1:
            raise RangeError(f"range bounds must be 1-indexed: {self!r}")

    @classmethod
    def parse(cls, text: str) -> Rect:
        """`'B5:E25'` or `'D14'` into a rectangle.

        A single cell becomes a 1×1 rectangle, so callers never branch on which form they were
        given — the branch is where an off-by-one lives.
        """
        stripped = text.strip().upper()
        if not _RANGE.match(stripped):
            raise RangeError(
                f"not an A1 range: {text!r}. Expected a form like 'B5' or 'B5:E25' - no sheet "
                "prefix, no '$' markers, no named range."
            )
        min_col, min_row, max_col, max_row = range_boundaries(stripped)
        # `range_boundaries` types these as optional because it also accepts open-ended forms like
        # 'B:E'; the regex above has already excluded those, so this is a narrowing, not a guess.
        if None in (min_col, min_row, max_col, max_row):  # pragma: no cover - excluded by the regex
            raise RangeError(f"open-ended range is not accepted: {text!r}")
        return cls(
            min_col=int(min_col),  # type: ignore[arg-type]
            min_row=int(min_row),  # type: ignore[arg-type]
            max_col=int(max_col),  # type: ignore[arg-type]
            max_row=int(max_row),  # type: ignore[arg-type]
        )

    @property
    def rows(self) -> int:
        return self.max_row - self.min_row + 1

    @property
    def cols(self) -> int:
        return self.max_col - self.min_col + 1

    @property
    def is_single_row(self) -> bool:
        return self.rows == 1

    @property
    def is_single_column(self) -> bool:
        return self.cols == 1

    def cell(self, row: int, col: int) -> str:
        """The A1 reference of one cell, in absolute sheet coordinates."""
        return f"{get_column_letter(col)}{row}"

    def cells_by_row(self) -> list[list[str]]:
        """Every cell reference, row-major. The shape a reader expects a table to have."""
        return [
            [self.cell(row, col) for col in range(self.min_col, self.max_col + 1)]
            for row in range(self.min_row, self.max_row + 1)
        ]

    def row_cells(self, row: int) -> list[str]:
        return [self.cell(row, col) for col in range(self.min_col, self.max_col + 1)]

    def column_cells(self, col: int) -> list[str]:
        return [self.cell(row, col) for row in range(self.min_row, self.max_row + 1)]

    def contains(self, other: Rect) -> bool:
        return (
            other.min_col >= self.min_col
            and other.max_col <= self.max_col
            and other.min_row >= self.min_row
            and other.max_row <= self.max_row
        )

    @property
    def rendered(self) -> str:
        """Back to A1, so a finding can quote the range it was given."""
        start = self.cell(self.min_row, self.min_col)
        if self.rows == 1 and self.cols == 1:
            return start
        return f"{start}:{self.cell(self.max_row, self.max_col)}"

    def __str__(self) -> str:
        return self.rendered
