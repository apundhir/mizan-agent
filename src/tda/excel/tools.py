"""The mapping agent's tool surface: it cannot return a value, by construction.

This module is the clearest demonstration in the repository of a distinction the whole design rests
on. `"Do not read the numbers"` in a prompt is a **request**. A function that has no code path
capable of putting a number in its return value is a **guarantee**. Only one of those survives a
model that is having an off day, and only one of them is worth telling the regulator about.

So: the agent decides *which sheet and which range hold which metric*, and it makes that decision
from the output of the two functions below. Neither can emit a cell's value. The second stage then
reads the values with `openpyxl` and no model involved (`tda.excel.read`).

## Why there is no tool-call loop

`ModelRequest` carries text, not tool calls — see `tda.agents.provider.base`. The two functions here
are therefore called by *code*, before the request is built, and their output is rendered into the
message. That is a stronger guarantee than a tool-use loop, not a weaker one: with tool use the model
holds a channel it can keep asking down, and each answer has to be filtered correctly every time.
Here the model has no channel at all. What `digest()` renders is everything it will ever see about
the workbook, and it is the same bytes the cassette key hashes, which is also what makes replay exact.

## Why redaction is semantic rather than type-based

The obvious implementation is "return the cell only if `isinstance(value, str)`". It is wrong, and
wrong in exactly the direction that matters: a percentage stored as text — `'81.70%'`, which
`Claim.raw_text` exists precisely to preserve, and which the story requires the parser to handle — is
a `str`. A type-based filter would hand the model the one class of value the design is built to keep
away from it, and would look completely correct in review.

So a cell is redacted when its *content* reads as a number, whatever its storage type. `'81.70%'`,
`'1,285'`, `'(44)'` and `'AED 612.40'` are all values. `'Occupancy %'`, `'January 2026'` and
`'2026-Q1 total'` are all labels, and the agent needs every one of them to do its job.

The boundary is deliberately drawn on the safe side. A header cell containing nothing but a bare year
— `'2026'` — is redacted, because no rule distinguishes it from a quantity without knowing what the
column means, which is the question being asked. The cost is that such a header is unmappable and
gets flagged for a human; the cost of the other default is a value in a prompt.

Formulas are redacted too, and for the same reason rather than a different one: `'=4140'` is a
literal wearing a disguise, and `data_only=False` returns the formula text.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from typing import TYPE_CHECKING, Final

if TYPE_CHECKING:
    from openpyxl.workbook.workbook import Workbook

# What the model sees in place of a value. Distinct tokens because the two cases mean different
# things to a reader of a recorded cassette: one is a number, the other is a computation.
REDACTED_VALUE: Final = "<value>"
REDACTED_FORMULA: Final = "<formula>"

# How much of a sheet `peek_headers` shows, and the rule is not a simple window.
#
# The agent needs two different things, and one window cannot supply both. It needs the **header
# block** — the top rows, where the labels that say what each column means live — and it needs the
# **extent of the label axis**, because a range it cannot see the end of is a range it has to guess
# at. A first attempt showed the top eight rows and produced a mapping agent that could see
# `Australia`, `China`, `Czech Republic`, `Egypt` and had no way to know the countries ran to row 26.
#
# So: every *label* in the used range is shown (labels are what the agent reasons about, and by
# nature they are not values), and redaction tokens are shown only for the first few rows, which is
# all it takes to see which columns hold values and where the block starts. The prompt then grows
# with the number of labels rather than with the size of the sheet.
VALUE_PREVIEW_ROWS: Final = 8

# A visible cap. A sheet with more labels than this is pathological, and the peek says so rather than
# quietly stopping — a truncation the model cannot see is a truncation that becomes a wrong mapping.
MAX_LABELS: Final = 200

# Decoration that surrounds a number without changing the fact that it is one. Stripped before the
# numeric test, so `'AED 612.40'` and `'81.70%'` are recognised as values.
_CURRENCY: Final = frozenset("$€£¥₹")
_UNIT_WORDS: Final = frozenset({"aed", "usd", "eur", "gbp", "pct", "percent"})

# A bare year, or any run of digits, standing alone. See the module docstring on why this is
# redacted rather than shown.
_DIGITS_ONLY: Final = re.compile(r"^[0-9]+$")


def looks_numeric(text: str) -> bool:
    """Whether a cell's text content reads as a quantity rather than as a label.

    The single decision this module turns on. Erring towards `True` withholds a label the agent
    could have used and costs a human mapping; erring towards `False` puts a value in a prompt. The
    second is the failure this whole module exists to prevent, so the tie goes to `True`.
    """
    stripped = text.strip()
    if not stripped:
        return False

    # A parenthesised negative, as accountants write it.
    if stripped.startswith("(") and stripped.endswith(")"):
        stripped = stripped[1:-1].strip()

    # Trailing or leading units and symbols: 'AED 612.40', '81.70%', '$1,285'.
    stripped = stripped.removesuffix("%").strip()
    stripped = "".join(ch for ch in stripped if ch not in _CURRENCY).strip()
    for word in _UNIT_WORDS:
        if stripped.lower().startswith(word):
            stripped = stripped[len(word) :].strip()
        if stripped.lower().endswith(word):
            stripped = stripped[: -len(word)].strip()

    if not stripped:
        # The cell held nothing but a currency symbol or a percent sign. Not a number, and not a
        # label either; treated as a label so the agent can see the sheet is oddly formed.
        return False

    if _DIGITS_ONLY.match(stripped):
        return True

    # Thousands separators, but only between digits - `'Q1,2026'` is not a number.
    candidate = re.sub(r"(?<=[0-9]),(?=[0-9]{3}\b)", "", stripped)
    try:
        Decimal(candidate)
    except (InvalidOperation, ValueError):
        return False
    return True


def render_cell(value: object) -> str | None:
    """One cell as the agent may see it, or `None` for a cell it is not shown at all.

    `None` means empty. Every non-empty cell yields either a label or a redaction token, and there
    is no third branch — which is what makes the guarantee in the module docstring checkable by
    reading this function rather than by trusting the caller.
    """
    if value is None:
        return None

    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        if text.startswith("="):
            return REDACTED_FORMULA
        return REDACTED_VALUE if looks_numeric(text) else text

    # Everything else is a quantity: int, float, Decimal, bool, and the date/time family, which
    # openpyxl returns as Python objects and which are ordinal numbers in the file.
    if isinstance(value, bool | int | float | Decimal | datetime | date | time | timedelta):
        return REDACTED_VALUE

    # An unanticipated type. Redacted rather than coerced to `str`: an unknown object's `repr` is
    # exactly where a value would slip through, and this branch is the one a future openpyxl
    # version could start using.
    return REDACTED_VALUE


@dataclass(frozen=True, slots=True)
class SheetSummary:
    """One sheet, as `list_sheets` reports it."""

    name: str
    used_range: str

    def render(self) -> str:
        return f"- {self.name!r} (used range {self.used_range})"


@dataclass(frozen=True, slots=True)
class PeekedCell:
    """One cell's A1 reference and what the agent is shown for it."""

    cell: str
    text: str


@dataclass(frozen=True, slots=True)
class SheetPeek:
    """One sheet's labels, plus enough redaction tokens to show where the values begin."""

    name: str
    used_range: str
    cells: tuple[PeekedCell, ...]
    labels_omitted: int = 0

    def render(self) -> str:
        lines = [f"## sheet {self.name!r} (used range {self.used_range})"]
        lines.extend(f"  {cell.cell}: {cell.text}" for cell in self.cells)
        if not self.cells:
            lines.append("  (the sheet is empty)")
        if self.labels_omitted:
            # Stated, not silent. A mapping made against a range whose end the agent could not see
            # is a guess, and it must be able to tell that is what it is being asked for.
            lines.append(
                f"  ... {self.labels_omitted} further label(s) not shown (cap {MAX_LABELS}). "
                "Do not map a range whose extent is not visible here."
            )
        return "\n".join(lines)


def list_sheets(workbook: Workbook) -> tuple[SheetSummary, ...]:
    """Every sheet and its used range. Names and geometry, never contents."""
    return tuple(
        SheetSummary(name=sheet.title, used_range=str(sheet.dimensions))
        for sheet in workbook.worksheets
    )


def peek_headers(
    workbook: Workbook, sheet: str, *, value_preview_rows: int = VALUE_PREVIEW_ROWS
) -> SheetPeek:
    """Every label in one sheet, and redaction tokens for the first few rows only.

    Raises `KeyError` for a sheet that does not exist rather than returning an empty peek: an empty
    peek and a missing sheet would look identical to the caller, and the second is a bug.

    Iteration is bounded by the sheet's *existing* extent, captured before the first access. This is
    not defensive tidiness — `iter_rows` beyond the used range **creates** the cells it walks, which
    grows `worksheet.dimensions`. An earlier version peeked a fixed 16 columns and made `Summary`
    report its used range as `A1:B11` in one part of the prompt and `A1:P11` in the next. Reading a
    workbook must not change it, and a prompt must not contradict itself.
    """
    if sheet not in workbook.sheetnames:
        raise KeyError(f"no sheet named {sheet!r}; have {workbook.sheetnames}")

    worksheet = workbook[sheet]
    used_range = str(worksheet.dimensions)
    max_row, max_col = worksheet.max_row, worksheet.max_column

    cells: list[PeekedCell] = []
    omitted = 0
    for row in worksheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
        for cell in row:
            text = render_cell(cell.value)
            if text is None:
                continue
            is_redacted = text in (REDACTED_VALUE, REDACTED_FORMULA)
            # `cell.row` is typed optional by openpyxl; a cell yielded by `iter_rows` always has
            # one, and treating an absent row as "past the preview" keeps the value hidden either
            # way - the safe direction for this particular default.
            if is_redacted and (cell.row or value_preview_rows + 1) > value_preview_rows:
                continue
            if not is_redacted and len(cells) >= MAX_LABELS:
                omitted += 1
                continue
            cells.append(PeekedCell(cell=str(cell.coordinate), text=text))

    return SheetPeek(
        name=worksheet.title,
        used_range=used_range,
        cells=tuple(cells),
        labels_omitted=omitted,
    )


def digest(workbook: Workbook) -> str:
    """Everything the mapping agent will ever see about this workbook, as one block of text.

    Assembled here rather than in the agent runner so that the guarantee and the thing being
    guaranteed sit in the same file. If a future change wants to give the model more context, it has
    to be made in this module, next to the docstring explaining why it must not be a value.
    """
    summaries = list_sheets(workbook)
    parts = ["# sheets", *(summary.render() for summary in summaries), "", "# labels"]
    parts.extend(peek_headers(workbook, summary.name).render() for summary in summaries)
    parts.append("")
    parts.append(
        f"Cells shown as {REDACTED_VALUE} hold a quantity and {REDACTED_FORMULA} holds a formula. "
        "Neither can be read at this stage, by design - map the ranges and the second stage will "
        "read them."
    )
    return "\n".join(parts)
