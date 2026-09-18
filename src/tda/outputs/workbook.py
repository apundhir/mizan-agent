"""The annotated workbook — the artefact the officer actually reads.

A hotel's submission comes back with every figure it claimed marked against what the reservation
records say: green verified, amber a policy difference, red a material variance with the computed
value and the page it came from, blue something that could not be checked at all.

## The original is never modified in place, and the code proves it rather than promising it

It is the hotel's submission and a piece of evidence. Annotating a copy means the comparison can
always be re-run against the untouched original, and means nobody can ask whether the tool changed
the thing it was judging.

That is easy to say and easy to break — one `load_workbook(original)` followed by a careless
`save()` and the evidence is gone with no error anywhere. So `annotate` digests the original before
and after and **raises if the two differ**. The check costs one hash of a file the ledger already
hashes, and it turns a promise into a postcondition.

## Every claimed cell is coloured, including the ones with nothing wrong

Colouring only the problems would leave a reader unable to distinguish *checked and correct* from
*not checked at all* — and D-SCOPE-02 is explicit that silence must not read as approval. Green
means "this figure was recomputed from the reservation records and matched". A cell with no colour
was not checked, and that now means something.

## Five colours, not four

PRD-92 names four. The fifth is grey, for a finding that is neither definitional nor material — a
V6 rounding difference, say. It exists because the alternative is worse in both directions: green
would tell a reader a cell verified clean when it produced a finding, and red would escalate a
rounding artefact into a material variance. A colour that overstates is how a correct system
produces an incorrect letter.

## Cells are found by the real reference; comments are redacted

The one place this file departs from "everything is rendered from the redacted verdict", and it is
forced. `tda.obs.redact` rewrites a sheet name carrying an honorific — `Dr. Ahmed Occupancy` is a
legal Excel sheet name — and a redacted reference finds no sheet. An earlier version looked cells up
through the redacted verdict and therefore drew *nothing* for such a finding, while the claim
underneath it stayed green: a material variance, coloured green, commented "the reservation records
agree". The redaction inverted the finding.

So the references are taken from the verdict as it was, and the **comment text** is redacted on the
way into the cell. That is consistent with why this artefact is exempt in the first place — it is a
copy of the hotel's own workbook going back to the hotel — and it keeps a citation pointing at
something. `verdict.json` still redacts the sheet name, and the cost of that is stated in ADR-0007:
a citation reading `[redacted:titled_name]!B5` cites nothing, which is one more reason for the
runbook's advice to ask the property to keep contact details out of the sheet.

## Findings that have no cell to sit on

A V7 extraction limit and a V5 missing claim carry `NotReached` as their Excel citation, by
construction — there is no cell, because the hotel wrote nothing there or the run never got that
far. They cannot be drawn, and a workbook that silently omitted them would be the most misleading
artefact in the set: all green, with the blocking findings invisible.

They are listed on the legend sheet and returned in `AnnotatedWorkbook.unplaced`, so the caller can
say so out loud — **and so is every other finding that could not be drawn**, including one naming a
sheet this workbook does not have. An earlier version dropped those silently and then printed
"Every finding in this verdict is marked on a cell" underneath, which is worse than omitting them:
an incomplete artefact that positively asserts it is complete.

## Two openpyxl properties worth knowing before reading one

**Cached formula values are lost.** openpyxl rewrites the file from its own model, and a formula's
last-computed value is not part of that model. Excel recomputes on open, so a reader sees the right
number; a script reading the annotated copy with `data_only=True` sees `None` where the original had
a cached value. The original is the file to read for values — which is the right default anyway.

**Comments are anchored to cells, not to a thread.** They are the old-style note, not a modern
threaded comment, and they are attributed to `Mizan` so nobody mistakes one for a colleague's.
"""

from __future__ import annotations

import shutil
from dataclasses import dataclass
from typing import TYPE_CHECKING, Final

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

from tda.contracts import ExcelRef, Severity, VarianceClass
from tda.obs.ledger import file_digest
from tda.obs.redact import redact
from tda.outputs.ooxml import make_reproducible

if TYPE_CHECKING:
    from collections.abc import Iterable, Sequence
    from pathlib import Path

    from openpyxl.workbook.workbook import Workbook
    from openpyxl.worksheet.worksheet import Worksheet

    from tda.contracts import Claim, Finding, Verdict

AUTHOR: Final = "Mizan"
LEGEND_SHEET: Final = "Mizan verification"

# Excel refuses to open a workbook with a longer note and offers to repair it instead. The one
# field that can reach this length is `Finding.narrative`, which is model prose and bounded
# nowhere upstream - so it is bounded here rather than at the point somebody's file will not open.
COMMENT_LIMIT: Final = 32_000

# Solid fills, light enough that the cell's own text stays readable in both Excel's light and dark
# renderings. A saturated fill on a numeric cell is a cell nobody can read, which defeats the point.
VERIFIED: Final = PatternFill("solid", fgColor="C6EFCE")
DEFINITIONAL: Final = PatternFill("solid", fgColor="FFE699")
MATERIAL: Final = PatternFill("solid", fgColor="FFC7CE")
BLOCKING: Final = PatternFill("solid", fgColor="BDD7EE")
OTHER: Final = PatternFill("solid", fgColor="D9D9D9")

LEGEND: Final = (
    ("Verified", VERIFIED, "Recomputed from the reservation records and matched."),
    (
        "Definitional (V2)",
        DEFINITIONAL,
        "A policy difference, not a hotel error. The comment names the rule that reproduces the "
        "figure exactly.",
    ),
    (
        "Material",
        MATERIAL,
        "The figure does not match what the records support. The comment carries the computed "
        "value and the page it came from.",
    ),
    ("Blocking", BLOCKING, "Something could not be checked. The run did not verify everything."),
    (
        "Other finding",
        OTHER,
        "A finding that is neither definitional nor material - a rounding difference, say.",
    ),
    ("No colour", None, "Not checked. Silence is not approval - D-SCOPE-02."),
)


@dataclass(frozen=True, slots=True)
class Unplaced:
    """A finding that is in the verdict and not on the workbook, and the reason.

    The reason is the point. "No cell exists because the hotel wrote nothing there" and "this names
    a sheet the workbook does not have" are a correct refusal and a defect respectively, and a
    reader who cannot tell them apart has to open the code to find out which they are looking at.
    """

    finding_id: str
    reason: str

    def render(self) -> str:
        return f"{self.finding_id}: {self.reason}"


@dataclass(frozen=True, slots=True)
class AnnotatedWorkbook:
    """Where the annotated copy landed, and what could not be drawn on it."""

    path: Path
    cells_coloured: int
    unplaced: tuple[Unplaced, ...]
    duplicate_cells: tuple[str, ...] = ()

    @property
    def unplaced_ids(self) -> tuple[str, ...]:
        return tuple(item.finding_id for item in self.unplaced)

    def render(self) -> str:
        lines = [f"  annotated workbook: {self.path}  ({self.cells_coloured} cell(s) marked)"]
        if self.unplaced:
            lines.append(
                f"  {len(self.unplaced)} finding(s) could not be drawn and are listed on the "
                f"'{LEGEND_SHEET}' sheet: {', '.join(i.render() for i in self.unplaced)}"
            )
        if self.duplicate_cells:
            lines.append(
                f"  {len(self.duplicate_cells)} cell(s) carry more than one claim and show only "
                f"the last: {', '.join(self.duplicate_cells)}. A workbook asserting two figures in "
                "one cell contradicts itself."
            )
        return "\n".join(lines)


class OriginalModifiedError(RuntimeError):
    """The submitted workbook changed while it was being annotated.

    Its own class because there is exactly one correct response — stop, and do not hand anybody the
    output — and a caller that catches `RuntimeError` broadly should still not swallow this one by
    accident.
    """


def annotate(
    workbook: Path, verdict: Verdict, claims: Sequence[Claim], directory: Path
) -> AnnotatedWorkbook:
    """Write an annotated copy of the submitted workbook into a run's artifact directory.

    The original is opened only to be hashed and copied. Every read and write below happens on the
    copy, and `_check_untouched` is what makes that a fact rather than an intention.

    A failure part-way through removes the copy before re-raising. The alternative is a file named
    `annotated_<workbook>.xlsx` sitting in the run directory whose name asserts something nobody
    did — and an officer who opens it has no way to know that.
    """
    _refuse_to_write_beside_the_evidence(workbook, directory)
    directory.mkdir(parents=True, exist_ok=True)
    before = file_digest(workbook)

    destination = directory / f"annotated_{workbook.name}"
    shutil.copy2(workbook, destination)
    try:
        annotated = _draw(destination, verdict, claims)
    except Exception:
        destination.unlink(missing_ok=True)
        raise
    _check_untouched(workbook, before)
    return annotated


def _draw(destination: Path, verdict: Verdict, claims: Sequence[Claim]) -> AnnotatedWorkbook:
    book = load_workbook(destination)
    # Before the marking loop, not after. Removing it afterwards would delete a mark this run had
    # already counted, and `cells_coloured` would report a cell that is not there.
    if LEGEND_SHEET in book.sheetnames:
        del book[LEGEND_SHEET]

    marks, unplaced, duplicates = _marks_for(verdict, claims, frozenset(book.sheetnames))
    coloured = 0
    for (sheet, cell), mark in sorted(marks.items()):
        target = book[sheet][cell]
        target.fill = mark.fill
        if mark.comment:
            target.comment = Comment(mark.comment, AUTHOR, height=160, width=380)
        coloured += 1

    _write_legend(book, verdict, unplaced, duplicates)
    book.save(destination)
    # After the save, not before: openpyxl stamps `dcterms:modified` with the wall clock inside
    # `save()` and discards whatever the caller set. See `tda.outputs.ooxml`.
    make_reproducible(destination)
    return AnnotatedWorkbook(
        path=destination,
        cells_coloured=coloured,
        unplaced=unplaced,
        duplicate_cells=duplicates,
    )


def _refuse_to_write_beside_the_evidence(workbook: Path, directory: Path) -> None:
    """The annotated copy never lands in the submission directory.

    Not a live path — `mizan run` writes to `artifacts/<run_id>/` — but this is the one function
    whose entire premise is leaving the submission alone, and a caller that passed the submission
    directory would drop a second `.xlsx` beside the evidence for the next run's file discovery to
    find.
    """
    if directory.resolve() == workbook.resolve().parent:
        raise OriginalModifiedError(
            f"refusing to write an annotated copy into {directory}, which is the directory the "
            "submission was read from. The next run's file discovery would find it, and a "
            "submission directory that accumulates this system's own output is no longer evidence."
        )


def _check_untouched(workbook: Path, before: str) -> None:
    """The postcondition: the submitted workbook is byte-identical to what we were handed.

    Extracted so it can be tested directly. Under every reachable input it cannot fire — nothing
    here opens the original for writing — which is exactly why it is worth having and worth
    testing: the failure it guards against is one careless `save(original)` away, and it would be
    silent.
    """
    after = file_digest(workbook)
    if before != after:
        raise OriginalModifiedError(
            f"{workbook.name} changed while it was being annotated ({before} -> {after}). The "
            "submitted workbook is evidence: the comparison must stay re-runnable against it, and "
            "nobody should have to ask whether the tool altered what it was judging."
        )


@dataclass(frozen=True, slots=True)
class _Mark:
    fill: PatternFill
    comment: str


# What wins when two findings land on one cell. Higher is more urgent, and the order errs towards
# saying more rather than less: a material variance hidden under an amber "a policy difference, not
# a hotel error" understates, and for a fee decision that is the worse direction to be wrong in.
_URGENCY: Final[dict[str, int]] = {"blocking": 3, "material": 2, "definitional": 1, "other": 0}


def _marks_for(
    verdict: Verdict, claims: Iterable[Claim], sheets: frozenset[str]
) -> tuple[dict[tuple[str, str], _Mark], tuple[Unplaced, ...], tuple[str, ...]]:
    """One mark per cell, plus everything that could not be given one.

    Claims first and findings second, deliberately: a cell with a finding is never green, and doing
    it in this order means that holds without a special case — the finding simply lands on top.
    Findings are then applied least-urgent first, so the most urgent is the one left showing.
    """
    marks: dict[tuple[str, str], _Mark] = {}
    duplicates: list[str] = []
    for claim in claims:
        where = (claim.excel_ref.sheet, claim.excel_ref.cell)
        if where in marks:
            # Two claims in one cell means the workbook asserts two figures in one place. If they
            # agree it is redundant and if they disagree it contradicts itself - either way the
            # reader must see it, which is the argument `tda.contracts.claim.index_by_key` already
            # makes about duplicate keys.
            duplicates.append(claim.excel_ref.citation)
        marks[where] = _Mark(VERIFIED, _comment(_verified_comment(claim)))

    unplaced: list[Unplaced] = []
    for finding in sorted((*verdict.findings, *verdict.definitional_items), key=_urgency_of):
        if not isinstance(finding.excel_ref, ExcelRef):
            unplaced.append(
                Unplaced(
                    finding.finding_id,
                    "no cell exists - the workbook states no figure for it, or the run stopped "
                    "before reading it",
                )
            )
            continue
        if finding.excel_ref.sheet not in sheets:
            # A claim-parse defect rather than a drawing problem, and it is reported rather than
            # dropped: a workbook that omitted the finding *and* printed "every finding is marked
            # on a cell" underneath is an incomplete artefact asserting it is complete.
            unplaced.append(
                Unplaced(
                    finding.finding_id,
                    f"names sheet {finding.excel_ref.sheet!r}, which this workbook does not have",
                )
            )
            continue
        marks[(finding.excel_ref.sheet, finding.excel_ref.cell)] = _Mark(
            _fill_for(finding), _comment(_finding_comment(finding))
        )
    return marks, tuple(unplaced), tuple(duplicates)


def _urgency_of(finding: Finding) -> int:
    if finding.severity is Severity.BLOCKING:
        return _URGENCY["blocking"]
    if finding.variance_class is VarianceClass.DEFINITIONAL:
        return _URGENCY["definitional"]
    if finding.severity is Severity.MATERIAL:
        return _URGENCY["material"]
    return _URGENCY["other"]


def _comment(text: str) -> str:
    """A cell note: redacted, and short enough that Excel will open the file.

    Redacted here rather than by rendering from the written verdict, for the reason the module
    docstring gives — the cell *references* have to be the real ones, or a redacted sheet name
    draws nothing and leaves the claim underneath it green.
    """
    cleaned, _ = redact(text)
    if len(cleaned) <= COMMENT_LIMIT:
        return cleaned
    return f"{cleaned[:COMMENT_LIMIT]}… (truncated; the full text is in verdict.json)"


def _fill_for(finding: Finding) -> PatternFill:
    """Severity first, then class. A blocking finding is blue whatever caused it: "this could not
    be checked" is the more urgent thing to say about a cell than what kind of variance it was."""
    if finding.severity is Severity.BLOCKING:
        return BLOCKING
    if finding.variance_class is VarianceClass.DEFINITIONAL:
        return DEFINITIONAL
    if finding.severity is Severity.MATERIAL:
        return MATERIAL
    return OTHER


def _verified_comment(claim: Claim) -> str:
    return (
        f"Verified.\n{claim.key.rendered}\n"
        f"Claimed {claim.value}, and the reservation records agree."
    )


def _finding_comment(finding: Finding) -> str:
    """What a reviewer needs in the cell itself, without opening another file.

    The per-class extras are the point. A material finding carries the computed value and the page
    it came from, because the next thing anybody does is check it. A definitional one names the
    rule that reproduces the claim exactly, because without that the amber reads as an accusation.
    """
    lines = [
        f"{finding.finding_id} · {finding.variance_class.value} · {finding.severity.value}",
        finding.key.rendered,
        f"Claimed: {_figure(finding.claimed)}",
        f"Computed: {_figure(finding.computed)}",
        f"Difference: {_figure(finding.difference)}",
    ]
    if finding.variance_class is VarianceClass.DEFINITIONAL:
        lines.append(f"Explained by: {finding.explaining_permutation}")
        if finding.also_explained_by:
            # Naming one cause as certain when several fit the evidence is a false precision that
            # a reviewer cannot see from the amber alone - D-CLS-09.
            lines.append(f"Also fits: {', '.join(finding.also_explained_by)}")
        lines.append("A policy difference, not a hotel error (D-MAT-06).")
    if finding.proposed_correction is not None:
        lines.append(f"Proposed: {finding.proposed_correction}")
    lines.append(f"Source: {finding.source_ref.citation}")
    lines.append(f"Clause: {finding.clause}")
    if finding.narrative:
        lines.append(finding.narrative)
    return "\n".join(lines)


def _figure(value: object) -> str:
    """`-` rather than `None`. A cell comment reading "Computed: None" looks like a defect in the
    tool; a dash reads as the absence it is."""
    return "-" if value is None else str(value)


def _write_legend(
    book: Workbook,
    verdict: Verdict,
    unplaced: tuple[Unplaced, ...],
    duplicates: tuple[str, ...] = (),
) -> None:
    """A first sheet carrying the verdict, the colour key and the findings with no cell.

    Added to the copy, never to the original. A workbook of coloured cells with no key is a puzzle
    rather than a report — and this is where the findings that could not be drawn get said out
    loud, which is the difference between an artefact that is incomplete and one that is
    misleading.
    """
    if LEGEND_SHEET in book.sheetnames:
        del book[LEGEND_SHEET]
    sheet: Worksheet = book.create_sheet(LEGEND_SHEET, 0)
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 96

    rows: list[tuple[str, str]] = [
        ("Mizan verification", ""),
        ("Verdict", verdict.status.value),
        ("Property", verdict.hotel_id),
        ("Period", verdict.period),
        ("Claims checked", str(verdict.claims_checked)),
        ("Hotel errors", str(verdict.hotel_error_count)),
        ("Definitional items", str(len(verdict.definitional_items))),
        ("Policy version", verdict.policy_version),
        ("Metric library", verdict.metric_library_version),
        ("Run", verdict.run_id),
        ("", ""),
    ]
    for label, value in rows:
        sheet.append([label, value])
    sheet["A1"].font = Font(bold=True, size=14)

    sheet.append(["Colour key", ""])
    sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)
    for name, fill, meaning in LEGEND:
        sheet.append([name, meaning])
        if fill is not None:
            sheet.cell(row=sheet.max_row, column=1).fill = fill
        sheet.cell(row=sheet.max_row, column=2).alignment = Alignment(
            wrap_text=True, vertical="top"
        )

    if duplicates:
        sheet.append(["", ""])
        sheet.append(["Cells claimed twice", ""])
        sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)
        for citation in duplicates:
            sheet.append(
                [citation, "More than one claim was read from this cell; only the last is shown."]
            )

    sheet.append(["", ""])
    sheet.append(["Findings not marked on a cell", ""])
    sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)
    if not unplaced:
        sheet.append(["None", "Every finding in this verdict is marked on a cell."])
        return
    sheet.append(
        [
            "",
            "These are in the verdict and are not drawn anywhere in this workbook. Read them here "
            "and in verdict.json.",
        ]
    )
    sheet.cell(row=sheet.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    for item in unplaced:
        sheet.append([item.finding_id, item.reason])
        sheet.cell(row=sheet.max_row, column=2).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
