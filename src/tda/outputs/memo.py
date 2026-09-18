"""The memo — one page, front-loaded, for the supervisor who reads the first block and nothing else.

That sentence is the whole design. If the verdict is not in the first block, the memo has failed at
its job even if every subsequent page is correct — so the status, the property, the period and the
counts are the first thing on the page, before any explanation of what was checked or how.

## The order is an argument, not a layout preference

1. **The verdict and the counts.** What happened, in four numbers.
2. **The findings table.** What a hotel would be written to about: metric, period, claimed,
   computed, difference, cause, the cell it was read from and the page it was checked against.
3. **Definitional items, in their own section, under their own heading.** Never in the table above.
   A policy disagreement listed among clerical errors is a correct finding that reads as an
   accusation, and D-MAT-06 exists because that mistake is expensive in a way no recount fixes.
4. **The signature block**, naming the policy version the verdict was produced under.

## The signature block cannot claim a review that did not happen

PRD-92 asks for a block "naming the reviewer". At the moment a run finishes there is no reviewer —
the review gate is PRD-91, and `Verdict.review_records` is empty until a human decides something. A
memo that printed a name anyway would be a forged sign-off on the one page a supervisor actually
reads.

So the block names whoever is in `review_records` when there is one, and when there is not it says
**"No human review has been recorded"** above a blank signature line. The document is still useful —
it is the thing the reviewer signs — and it never asserts that somebody looked.

## One page, and what happens when it is not

A verdict with forty findings does not fit on a page and should not pretend to. The first block is
fixed-size and always fits; the table grows. That is the right way round: the part that must be read
is the part that cannot be pushed off the page.
"""

from __future__ import annotations

from typing import TYPE_CHECKING, Final

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from tda.contracts import ExcelRef, Severity
from tda.outputs.ooxml import make_reproducible
from tda.outputs.verdict import VerdictSummary

if TYPE_CHECKING:
    from pathlib import Path

    from docx.document import Document as DocxDocument

    from tda.contracts import Finding, Verdict

MEMO_FILE: Final = "memo.docx"

# The status colours a supervisor reads before the words. Muted rather than saturated: a memo is a
# document that gets printed and filed, not a dashboard.
STATUS_COLOURS: Final[dict[str, RGBColor]] = {
    "PASS": RGBColor(0x1E, 0x6B, 0x2E),
    "FAIL": RGBColor(0xA3, 0x1D, 0x1D),
    "ESCALATED": RGBColor(0xA3, 0x1D, 0x1D),
    "REJECTED": RGBColor(0x7A, 0x4A, 0x00),
    "HALTED": RGBColor(0x7A, 0x4A, 0x00),
}

FINDING_COLUMNS: Final = (
    "Finding",
    "Metric",
    "Period",
    "Claimed",
    "Computed",
    "Difference",
    "Cause",
    "Excel cell",
    "Source",
)


def write_memo(directory: Path, verdict: Verdict) -> Path:
    """Render the memo into a run's artifact directory and return where it landed."""
    directory.mkdir(parents=True, exist_ok=True)
    document = Document()
    _page_setup(document)

    _heading(document, verdict)
    _first_block(document, verdict)
    _findings_table(document, verdict)
    _definitional_section(document, verdict)
    _not_verifiable_section(document, verdict)
    _signature_block(document, verdict)

    path = directory / MEMO_FILE
    # `str`, because python-docx's stub takes a path string or a stream and not a `Path`.
    document.save(str(path))
    # A `.docx` is a zip, and a zip carries the clock in every entry. See `tda.outputs.ooxml`.
    make_reproducible(path)
    return path


def _page_setup(document: DocxDocument) -> None:
    section = document.sections[0]
    section.left_margin = section.right_margin = Inches(0.8)
    section.top_margin = section.bottom_margin = Inches(0.7)
    normal = document.styles["Normal"]
    normal.font.size = Pt(9.5)


def _heading(document: DocxDocument, verdict: Verdict) -> None:
    title = document.add_paragraph()
    run = title.add_run("Quarterly verification memo")
    run.bold = True
    run.font.size = Pt(15)

    subtitle = document.add_paragraph()
    subtitle_run = subtitle.add_run(f"{verdict.hotel_id} · {verdict.period} · run {verdict.run_id}")
    subtitle_run.font.size = Pt(9)
    subtitle_run.italic = True


def _first_block(document: DocxDocument, verdict: Verdict) -> None:
    """The verdict and the counts. Everything below this is detail a supervisor may not reach."""
    summary = VerdictSummary.of(verdict)

    line = document.add_paragraph()
    label = line.add_run(verdict.status.value)
    label.bold = True
    label.font.size = Pt(20)
    label.font.color.rgb = STATUS_COLOURS.get(verdict.status.value, RGBColor(0, 0, 0))
    if verdict.rejection_reason:
        reason = line.add_run(f"  {verdict.rejection_reason.value.replace('_', ' ')}")
        reason.font.size = Pt(11)

    table = document.add_table(rows=2, cols=4)
    table.style = "Table Grid"
    headers = ("Claims checked", "Hotel errors", "Definitional items", "Not verifiable")
    figures = (
        str(verdict.claims_checked),
        str(summary.hotel_errors),
        str(summary.definitional_items),
        str(summary.not_verifiable),
    )
    for column, (header, figure) in enumerate(zip(headers, figures, strict=True)):
        head = table.cell(0, column).paragraphs[0]
        head_run = head.add_run(header)
        head_run.bold = True
        head_run.font.size = Pt(8)
        body = table.cell(1, column).paragraphs[0]
        body_run = body.add_run(figure)
        body_run.font.size = Pt(14)

    if blocking := summary.by_severity.get(Severity.BLOCKING.value, 0):
        warning = document.add_paragraph()
        warning_run = warning.add_run(
            f"{blocking} blocking finding(s): this run did not verify everything it was given."
        )
        warning_run.bold = True

    provenance = document.add_paragraph()
    provenance_run = provenance.add_run(
        f"Policy {verdict.policy_version} · metric library {verdict.metric_library_version} · "
        f"model {verdict.model_id} ({verdict.provider_mode})"
    )
    provenance_run.font.size = Pt(8)
    provenance_run.italic = True


def _findings_table(document: DocxDocument, verdict: Verdict) -> None:
    _section_heading(document, "Findings")
    if not verdict.findings:
        document.add_paragraph(
            "None. Every figure the workbook states was recomputed from the reservation records "
            "and matched."
        )
        return

    table = document.add_table(rows=1, cols=len(FINDING_COLUMNS))
    table.style = "Table Grid"
    for column, name in enumerate(FINDING_COLUMNS):
        run = table.cell(0, column).paragraphs[0].add_run(name)
        run.bold = True
        run.font.size = Pt(7.5)

    for finding in verdict.findings:
        _finding_row(table.add_row().cells, finding)


def _definitional_section(document: DocxDocument, verdict: Verdict) -> None:
    """Its own heading, its own table, and a sentence saying what it is not.

    The sentence is not decoration. Without it the section reads as a second list of things the
    hotel got wrong, and the whole point of separating them is that they are not.
    """
    _section_heading(document, "Definitional items — not hotel errors")
    if not verdict.definitional_items:
        document.add_paragraph("None.")
        return

    note = document.add_paragraph()
    note_run = note.add_run(
        "These figures are reproduced exactly by an alternative reading of the definitions. They "
        "are a question for the policy owner, not a correction for the property, and they are "
        "excluded from the hotel error count above (D-MAT-06)."
    )
    note_run.italic = True

    table = document.add_table(rows=1, cols=len(FINDING_COLUMNS))
    table.style = "Table Grid"
    for column, name in enumerate(FINDING_COLUMNS):
        run = (
            table.cell(0, column).paragraphs[0].add_run("Explained by" if name == "Cause" else name)
        )
        run.bold = True
        run.font.size = Pt(7.5)

    for finding in verdict.definitional_items:
        _finding_row(table.add_row().cells, finding, cause=_explanation(finding))


def _not_verifiable_section(document: DocxDocument, verdict: Verdict) -> None:
    """What nobody could check, and what is missing to check it.

    Included because a memo listing only what was checked lets silence read as approval, which
    D-SCOPE-02 forbids in the verdict and which would be no less wrong on paper.
    """
    if not verdict.not_verifiable and not verdict.out_of_scope_claims:
        return
    _section_heading(document, "Not verified")
    for item in verdict.not_verifiable:
        # Both halves. `reason` is the code and `detail` is what a reviewer needs in order
        # to make it verifiable next time - printing only the first leaves the reader with
        # a diagnosis and no remedy.
        document.add_paragraph(
            f"{item.key.rendered}: {item.reason} — {item.detail}", style="List Bullet"
        )
    for key in verdict.out_of_scope_claims:
        document.add_paragraph(
            f"{key.rendered}: out of scope for this verification",
            style="List Bullet",
        )


def _signature_block(document: DocxDocument, verdict: Verdict) -> None:
    _section_heading(document, "Review and signature")
    applied = document.add_paragraph()
    applied_run = applied.add_run(
        f"Policy version applied: {verdict.policy_version}  ·  "
        f"Metric library: {verdict.metric_library_version}"
    )
    applied_run.bold = True

    if verdict.review_records:
        # The standing decision per finding, not every record. A supervisor reading two
        # contradictory lines for one finding has to work out which one is in force, and the
        # answer is already known - `Verdict.standing_decisions` knows it. The earlier decisions
        # stay in `verdict.json`, where an auditor looks, and a changed mind is noted rather than
        # hidden.
        for finding_id, record in sorted(verdict.standing_decisions.items()):
            earlier = sum(1 for r in verdict.review_records if r.finding_id == finding_id) - 1
            changed = f"  (superseded {earlier} earlier decision(s))" if earlier else ""
            # The amended figure, spelled out. Without it the forwarded memo reads "amend by X"
            # and leaves the hotel exactly nothing to restate the return to - which is the
            # condition `record_decision` refuses an amendment for in the first place.
            amended = f" → {record.amended_value}" if record.amended_value else ""
            document.add_paragraph(
                f"{record.finding_id}: {record.decision.value}{amended} by {record.reviewer} on "
                f"{record.decided_at:%Y-%m-%d %H:%M %Z}"
                f"{f' — {record.note}' if record.note else ''}{changed}",
                style="List Bullet",
            )
        if undecided := verdict.undecided_findings:
            document.add_paragraph(
                f"Undecided: {', '.join(undecided)}. This verdict is not fully reviewed."
            )
    else:
        # See the module docstring. Printing a name nobody supplied would be a forged sign-off on
        # the one page a supervisor actually reads.
        pending = document.add_paragraph()
        pending_run = pending.add_run(
            "No human review has been recorded against this verdict. The figures above are the "
            "system's, and nothing on this page has been accepted by a person yet."
        )
        pending_run.italic = True

    document.add_paragraph()
    line = document.add_paragraph(
        "Reviewed by: ______________________________    Date: ____________"
    )
    line.alignment = WD_ALIGN_PARAGRAPH.LEFT


def _section_heading(document: DocxDocument, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(11)


def _finding_row(cells: list, finding: Finding, cause: str | None = None) -> None:  # type: ignore[type-arg]
    """One row of the findings table. Kept in one place so both tables cannot drift apart."""
    values = (
        finding.finding_id,
        finding.key.metric.value + (f" {finding.key.value}" if finding.key.value else ""),
        finding.key.period,
        _figure(finding.claimed),
        _figure(finding.computed),
        _figure(finding.difference),
        cause or f"{finding.variance_class.value} · {finding.clause}",
        finding.excel_ref.citation if isinstance(finding.excel_ref, ExcelRef) else "—",
        finding.source_ref.citation,
    )
    for column, value in enumerate(values):
        run = cells[column].paragraphs[0].add_run(value)
        run.font.size = Pt(7.5)


def _explanation(finding: Finding) -> str:
    """What reproduces a definitional claim, and honestly whether it is the only thing that does."""
    also = (
        f" (also fits {', '.join(finding.also_explained_by)})" if finding.also_explained_by else ""
    )
    return f"{finding.explaining_permutation}{also}"


def _figure(value: object) -> str:
    """An em dash, not `None`. A memo that prints "None" in a numeric column looks like a defect
    in the tool rather than the stated absence it is."""
    return "—" if value is None else str(value)
