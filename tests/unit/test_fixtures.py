"""The fixture producer, checked against the one property the whole story rests on.

PRD-94's acceptance criterion is that an expectation is **derived** and never written. That is a
claim about code, so most of this file is written to fail if the code stopped being like that:

- **Sensitivity.** Retarget the transposition at another country and the expected *cell* must move
  on its own. Change the permutation and the expected *value* must move. A derivation that returned
  a constant, or one that read the answer out of a table somebody maintained by hand, passes
  neither.
- **Fail-loud.** A target that names no claim, a control fixture that somehow produces a mismatch,
  a corpus that no longer matches its own manifest: each must raise rather than produce a smaller
  expectation. The quiet version of every one of those failures is a green eval.
- **Policy, not memory.** Severity and escalation are asserted to follow a *changed* policy
  document, not to equal the strings in the shipped one. Asserting the literal would pass equally
  well against a derivation that typed them in.

The two digest tests are the postconditions rather than assertions about behaviour: the demo corpus
must be byte-identical after a full build, and two builds must agree. Both are checked by deleting
the guarantee, which for the corpus guard means tampering with a copied corpus and watching the
check fire.
"""

from __future__ import annotations

import copy
import dataclasses
import json
import shutil
import zipfile
from decimal import Decimal
from pathlib import Path
from typing import TYPE_CHECKING, Any

import jsonschema
import pytest
from openpyxl import load_workbook

from datagen.aggregate import load_policy_view, occupancy_pct
from datagen.claims import baseline_table, by_key
from datagen.ledger import generate_inventory
from datagen.spec import (
    MONTHS,
    PINNED_OOXML_MODIFIED,
    PINNED_ZIP_DATE_TIME,
    QUARTER,
    WORKBOOK,
)
from fixtures import FIXTURE_SET_VERSION
from fixtures.__main__ import main
from fixtures.derive import (
    SCHEMA_PATH,
    _halt_at_first_blocking,
    _source_of,
    _status,
    expectation,
    load_classification_view,
)
from fixtures.materialise import (
    check_demo_untouched,
    demo_digests,
    load_truth,
    materialise_all,
    materialise_one,
)
from fixtures.mutate import (
    load_permutation_view,
    load_policy_document,
    mutated,
    resolve,
    transpose_digits,
)
from fixtures.spec import (
    FIXTURES,
    FixtureError,
    FixtureSpec,
    Mutation,
    MutationKind,
    Target,
    fixture,
)

if TYPE_CHECKING:
    from datagen.aggregate import PolicyView
    from datagen.claims import ClaimCell
    from fixtures.derive import ClassificationView
    from fixtures.materialise import FixtureBuild
    from fixtures.mutate import PermutationView

REPO_ROOT = Path(__file__).resolve().parents[2]
DEMO = REPO_ROOT / "corpus" / "demo"


# ── shared state, built once ─────────────────────────────────────────────────


@pytest.fixture(scope="module")
def truth() -> dict[str, float]:
    return load_truth(DEMO)


@pytest.fixture(scope="module")
def policy() -> PolicyView:
    return load_policy_view()


@pytest.fixture(scope="module")
def document() -> dict[str, object]:
    return load_policy_document()


@pytest.fixture(scope="module")
def permutations(document: dict[str, object]) -> PermutationView:
    return load_permutation_view(document)


@pytest.fixture(scope="module")
def view(document: dict[str, object]) -> ClassificationView:
    return load_classification_view(document)


@pytest.fixture(scope="module")
def table(truth: dict[str, float]) -> tuple[ClaimCell, ...]:
    return baseline_table(MONTHS, truth)


@pytest.fixture(scope="module")
def built(tmp_path_factory: pytest.TempPathFactory) -> tuple[FixtureBuild, ...]:
    """One full materialisation, shared. Building it three times would test the clock."""
    return materialise_all(DEMO, tmp_path_factory.mktemp("fixtures"))


def _derive(
    spec: FixtureSpec,
    table: tuple[ClaimCell, ...],
    truth: dict[str, float],
    policy: PolicyView,
    permutations: PermutationView,
    view: ClassificationView,
) -> dict[str, Any]:
    return expectation(spec, mutated(spec, table, policy, permutations), truth, policy, view)


def _payload(built: tuple[FixtureBuild, ...], fixture_id: str) -> dict[str, Any]:
    for build in built:
        if build.fixture_id == fixture_id:
            loaded = json.loads((build.path / "expected.json").read_text(encoding="utf-8"))
            assert isinstance(loaded, dict)
            return loaded
    raise AssertionError(f"{fixture_id} was not built")


# ── the corpus is never written to ───────────────────────────────────────────


def test_the_demo_corpus_is_byte_identical_after_a_full_materialisation(
    built: tuple[FixtureBuild, ...],
) -> None:
    """The point of the whole producer: it mutates a copy.

    `built` has already run a complete materialisation by the time this asserts, so the digests
    below are taken after three workbooks were rendered and fifteen files copied.
    """
    assert built, "nothing was built, so this would assert nothing"

    manifest = json.loads((DEMO / "manifest.json").read_text(encoding="utf-8"))
    assert demo_digests(DEMO) == dict(manifest["files"])


def test_the_corpus_guard_fires_when_a_demo_file_moves(tmp_path: Path) -> None:
    """The postcondition, watched failing. Against a copy, because the real one must not move.

    Without this the digest check is untested code: under every reachable input it cannot fire,
    which is exactly the argument `tda.outputs.workbook._check_untouched` makes for testing its
    own equivalent directly.
    """
    replica = tmp_path / "demo"
    shutil.copytree(DEMO, replica)
    before = demo_digests(replica)

    (replica / "submission" / "inventory_2026-Q1.csv").write_text("tampered\n", encoding="utf-8")

    with pytest.raises(FixtureError, match="changed while fixtures were being built"):
        check_demo_untouched(replica, before)


def test_a_corpus_that_does_not_match_its_manifest_is_refused(tmp_path: Path) -> None:
    """Building on a drifted corpus would derive expectations correct with respect to nothing."""
    replica = tmp_path / "demo"
    shutil.copytree(DEMO, replica)
    (replica / "ground_truth" / "truth_metrics.json").write_text("{}", encoding="utf-8")

    with pytest.raises(FixtureError, match="does not match its own manifest"):
        materialise_all(replica, tmp_path / "out")


# ── reproducibility ──────────────────────────────────────────────────────────


def test_materialising_twice_is_byte_identical(tmp_path: Path) -> None:
    first = materialise_all(DEMO, tmp_path / "one")
    second = materialise_all(DEMO, tmp_path / "two")

    assert [build.files for build in first] == [build.files for build in second]
    assert all(build.files for build in first), "a fixture with no files would pass vacuously"


@pytest.mark.parametrize("fixture_id", [spec.fixture_id for spec in FIXTURES])
def test_every_rendered_workbook_carries_the_pinned_timestamps(
    built: tuple[FixtureBuild, ...], fixture_id: str
) -> None:
    """The reproducibility fix, asserted structurally rather than by building twice.

    Two builds a second apart agree even without `finalise_xlsx`, because openpyxl stamps the wall
    clock and the wall clock has not moved yet. So the pair-of-builds test passes with the
    guarantee removed, which is the defect this project keeps finding: a test that still passes
    when the behaviour is deleted. What actually has to hold is that every zip entry carries the
    pinned date and `dcterms:modified` was rewritten after the save, and both are checkable in one
    run of one build.
    """
    build = next(b for b in built if b.fixture_id == fixture_id)
    workbook = build.path / "submission" / "claims_2026-Q1.xlsx"

    with zipfile.ZipFile(workbook) as archive:
        stamps = {info.date_time for info in archive.infolist()}
        core = archive.read("docProps/core.xml")

    assert stamps == {PINNED_ZIP_DATE_TIME}, (
        "a zip entry carries the time the build ran, so this fixture's digest changes on every "
        "rebuild and `--verify` becomes a clock comparison"
    )
    modified = core.split(b"<dcterms:modified")[1].split(b"</dcterms:modified>")[0]
    assert PINNED_OOXML_MODIFIED.encode() in modified


def test_verify_accepts_a_freshly_built_tree(tmp_path: Path) -> None:
    out = tmp_path / "fixtures"
    assert main(["--demo", str(DEMO), "--out", str(out)]) == 0
    assert main(["--demo", str(DEMO), "--out", str(out), "--verify"]) == 0


def test_verify_rejects_a_hand_edited_expectation(tmp_path: Path) -> None:
    """The failure this exists for: an expectation edited until the eval went green."""
    out = tmp_path / "fixtures"
    assert main(["--demo", str(DEMO), "--out", str(out)]) == 0

    path = out / "F2" / "expected.json"
    payload = json.loads(path.read_text(encoding="utf-8"))
    payload["findings"] = []
    payload["status"] = "PASS"
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")

    assert main(["--demo", str(DEMO), "--out", str(out), "--verify"]) == 1


def test_verify_rejects_a_hand_edited_workbook(tmp_path: Path) -> None:
    out = tmp_path / "fixtures"
    assert main(["--demo", str(DEMO), "--out", str(out)]) == 0

    workbook = out / "F3" / "submission" / "claims_2026-Q1.xlsx"
    workbook.write_bytes(workbook.read_bytes() + b"\n")

    assert main(["--demo", str(DEMO), "--out", str(out), "--verify"]) == 1


def test_verify_reports_a_missing_tree_rather_than_building_one(tmp_path: Path) -> None:
    assert main(["--demo", str(DEMO), "--out", str(tmp_path / "absent"), "--verify"]) == 1


def test_a_missing_corpus_is_a_bad_invocation(tmp_path: Path) -> None:
    assert main(["--demo", str(tmp_path / "nowhere"), "--out", str(tmp_path / "out")]) == 2


# ── the derivation is sensitive to the spec ──────────────────────────────────


def test_retargeting_the_transposition_moves_the_expected_cell(
    table: tuple[ClaimCell, ...],
    truth: dict[str, float],
    policy: PolicyView,
    permutations: PermutationView,
    view: ClassificationView,
) -> None:
    """The acceptance criterion, stated as a test.

    Nothing about the expectation is written down, so pointing the same mutation at France has to
    move the cell by itself. A derivation carrying a hand-maintained cell map passes the shipped
    F2 and fails here, which is the whole reason this test is not simply "F2 expects B10".
    """
    germany = _derive(fixture("F2"), table, truth, policy, permutations, view)
    france = _derive(
        FixtureSpec(
            fixture_id="F2",
            why=fixture("F2").why,
            mutation=Mutation(
                kind=MutationKind.TRANSPOSE_DIGITS,
                target=Target("guests_by_nationality", "2026-01", "FR"),
            ),
        ),
        table,
        truth,
        policy,
        permutations,
        view,
    )

    indexed = by_key(table)
    expected_cell = indexed["guests_by_nationality:2026-01:nationality_iso2=FR"].cell
    cells = {f["evidence"]["excel"]["cell"] for f in france["findings"]}

    assert expected_cell in cells
    assert cells != {f["evidence"]["excel"]["cell"] for f in germany["findings"]}
    assert {f["key"] for f in france["findings"]} == {
        "guests_by_nationality:2026-01:nationality_iso2=FR",
        "guests_by_nationality:2026-Q1:nationality_iso2=FR",
    }


def test_changing_the_permutation_moves_the_expected_value(
    table: tuple[ClaimCell, ...],
    truth: dict[str, float],
    policy: PolicyView,
    permutations: PermutationView,
    view: ClassificationView,
) -> None:
    """F3's claimed figure is computed from the permutation, not chosen to look wrong.

    `P-OOO-INCLUDED` leaves the denominator a sum of days and only puts the out-of-order rooms back
    into it, so February moves and the expected claimed value has to move with it. A derivation
    holding a literal 2141.67 would fail this and nothing else.
    """
    rooms = _derive(fixture("F3"), table, truth, policy, permutations, view)
    out_of_order = _derive(
        FixtureSpec(
            fixture_id="F3",
            why=fixture("F3").why,
            mutation=Mutation(
                kind=MutationKind.RECOMPUTE_UNDER,
                target=Target("occupancy_pct"),
                permutation="P-OOO-INCLUDED",
            ),
        ),
        table,
        truth,
        policy,
        permutations,
        view,
    )

    def claimed(payload: dict[str, Any], key: str) -> str | None:
        for item in payload["definitional_items"]:
            if item["key"] == key:
                claim = item["claimed"]
                assert isinstance(claim, str)
                return claim
        return None

    february = "occupancy_pct:2026-02"
    assert claimed(rooms, february) is not None
    assert claimed(out_of_order, february) is not None
    assert claimed(rooms, february) != claimed(out_of_order, february)
    assert {item["explaining_permutation"] for item in out_of_order["definitional_items"]} == {
        "P-OOO-INCLUDED"
    }


def test_the_definitional_value_is_computed_from_the_inventory_reference(
    table: tuple[ClaimCell, ...],
    truth: dict[str, float],
    policy: PolicyView,
    permutations: PermutationView,
    view: ClassificationView,
) -> None:
    """Recompute January by hand and require the fixture to agree.

    The denominator under `P-OCC-DENOM-ROOMS` is the room *count*, which is the maximum sellable
    rooms across the period, and the numerator is the hotel's own stated room-nights sold. Both
    halves are restated here from the inventory rather than read off the fixture, so a builder that
    divided by the wrong thing fails rather than agreeing with itself.
    """
    january = [row for row in generate_inventory() if row.day.month == 1]
    denominator = max(row.rooms_total - row.rooms_out_of_order for row in january)
    sold = int(by_key(table)["room_nights_sold:2026-01"].value)
    by_hand = occupancy_pct(sold, denominator, policy)

    payload = _derive(fixture("F3"), table, truth, policy, permutations, view)
    item = next(i for i in payload["definitional_items"] if i["key"] == "occupancy_pct:2026-01")

    assert denominator == 60, "the demo property has 60 rooms and no January out-of-order window"
    assert Decimal(item["claimed"]) == by_hand
    assert by_hand > 100, "occupancy is not capped, and the room-count denominator proves it"


def test_a_target_naming_no_claim_is_refused(
    table: tuple[ClaimCell, ...],
) -> None:
    """A spec that has drifted from the corpus must not materialise a clean workbook quietly.

    Two shapes of the same mistake: a dimension value the corpus does not carry, and a period it
    does not cover. Both would otherwise mutate nothing, derive nothing and pass beside F1.
    """
    with pytest.raises(FixtureError, match="names no claim"):
        resolve(table, Target("guests_by_nationality", "2026-01", "ZZ"))
    with pytest.raises(FixtureError, match="names no claim"):
        resolve(table, Target("occupancy_pct", "2025-12"))


def test_a_drifted_target_fails_the_whole_build(
    table: tuple[ClaimCell, ...], policy: PolicyView, permutations: PermutationView
) -> None:
    spec = FixtureSpec(
        fixture_id="F2",
        why=fixture("F2").why,
        mutation=Mutation(
            kind=MutationKind.TRANSPOSE_DIGITS,
            target=Target("guests_by_nationality", "2026-01", "ZZ"),
        ),
    )
    with pytest.raises(FixtureError, match="names no claim"):
        mutated(spec, table, policy, permutations)


def test_severity_and_escalation_follow_a_changed_policy(
    table: tuple[ClaimCell, ...],
    truth: dict[str, float],
    policy: PolicyView,
    permutations: PermutationView,
    document: dict[str, object],
) -> None:
    """Read from `policy.yaml`, not remembered.

    Asserting that a V1 is `material` and escalates to `hotel` would pass just as well against a
    derivation with those two strings typed into it. So the policy document is amended and the
    derived expectation has to follow it, which only a derivation that reads policy can do.
    """
    amended = copy.deepcopy(document)
    classes = amended["classification"]["classes"]  # type: ignore[index]
    classes["V1"]["escalates_to"] = "human_review"
    classes["V1"]["proposes_correction"] = False
    amended["materiality"]["severity_by_class"]["V1"] = "informational"  # type: ignore[index]

    payload = _derive(
        fixture("F2"), table, truth, policy, permutations, load_classification_view(amended)
    )

    assert {f["severity"] for f in payload["findings"]} == {"informational"}
    assert {f["escalates_to"] for f in payload["findings"]} == {"human_review"}
    assert {f["proposed_correction"] for f in payload["findings"]} == {None}


# ── the expectations themselves ──────────────────────────────────────────────


@pytest.mark.parametrize("fixture_id", [spec.fixture_id for spec in FIXTURES])
def test_every_expectation_on_disk_validates_against_the_schema(
    built: tuple[FixtureBuild, ...], fixture_id: str
) -> None:
    """Validated here with jsonschema directly rather than through `derive.validate`.

    Going through the producer's own validator would pass if that validator had quietly become a
    no-op, and the file on disk is what the scorer reads.
    """
    schema = json.loads(SCHEMA_PATH.read_text(encoding="utf-8"))
    jsonschema.Draft202012Validator.check_schema(schema)
    jsonschema.Draft202012Validator(schema).validate(_payload(built, fixture_id))


@pytest.mark.parametrize("fixture_id", [spec.fixture_id for spec in FIXTURES])
def test_every_expectation_carries_the_fixture_set_version(
    built: tuple[FixtureBuild, ...], fixture_id: str
) -> None:
    payload = _payload(built, fixture_id)
    assert payload["fixture_set_version"] == FIXTURE_SET_VERSION
    assert payload["exhaustive"] is True


def test_the_control_fixture_expects_nothing_at_all(built: tuple[FixtureBuild, ...]) -> None:
    """Precision, and the only fixture where an empty list is the right answer."""
    payload = _payload(built, "F1")

    assert payload["status"] == "PASS"
    assert payload["findings"] == []
    assert payload["definitional_items"] == []


def test_the_control_fixture_refuses_to_explain_a_mismatch(
    table: tuple[ClaimCell, ...],
    truth: dict[str, float],
    policy: PolicyView,
    permutations: PermutationView,
    view: ClassificationView,
) -> None:
    """A control that found something would mean the corpus and its truth had diverged.

    There is no class for it, so there must be no quiet answer either. The table handed in is F2's
    mutated one against F1's spec, which is exactly the shape that divergence would take.
    """
    damaged = mutated(fixture("F2"), table, policy, permutations)

    with pytest.raises(FixtureError, match="cannot produce a value mismatch"):
        expectation(fixture("F1"), damaged, truth, policy, view)


def test_the_transposition_expects_two_transcription_findings(
    built: tuple[FixtureBuild, ...],
) -> None:
    """The month **and** the quarter roll-up, because the roll-up is a claim in its own right.

    A producer that mutated only the cell it was pointed at would expect one finding here, and the
    quarter total carrying the same mistype would go unreported by a pipeline nobody had asked
    about it.
    """
    payload = _payload(built, "F2")

    assert payload["status"] == "FAIL"
    assert payload["definitional_items"] == []
    assert {f["key"] for f in payload["findings"]} == {
        "guests_by_nationality:2026-01:nationality_iso2=DE",
        "guests_by_nationality:2026-Q1:nationality_iso2=DE",
    }
    assert {f["variance_class"] for f in payload["findings"]} == {"V1"}
    assert {f["severity"] for f in payload["findings"]} == {"material"}
    assert {f["escalates_to"] for f in payload["findings"]} == {"hotel"}
    assert all(f["proposed_correction"] == f["computed"] for f in payload["findings"])
    assert all(f["explaining_permutation"] is None for f in payload["findings"])
    assert all(
        f["evidence"]["excel"]["sheet"] == WORKBOOK.nationality_sheet for f in payload["findings"]
    )


def test_the_stated_total_row_is_not_a_finding(built: tuple[FixtureBuild, ...]) -> None:
    """The `Total` row moves when a month does, and is never compared against anything.

    `tda.excel.read` treats it as a `StatedTotal`, so it has no metric key and can raise no
    finding. If the claim table ever started carrying it, this fixture would expect three findings
    and the third would be filed against a country called Total.
    """
    payload = _payload(built, "F2")
    assert len(payload["findings"]) == 2


def test_the_definitional_fixture_escalates_and_names_the_permutation(
    built: tuple[FixtureBuild, ...],
) -> None:
    """A definitional disagreement is not a hotel error, and must not be filed as one."""
    payload = _payload(built, "F3")

    assert payload["status"] == "ESCALATED"
    assert payload["findings"] == [], "a V2 in the findings array is a contract violation"
    assert {item["key"] for item in payload["definitional_items"]} == {
        f"occupancy_pct:{period}" for period in (*MONTHS, QUARTER)
    }
    assert {item["variance_class"] for item in payload["definitional_items"]} == {"V2"}
    assert {item["escalates_to"] for item in payload["definitional_items"]} == {"policy_owner"}
    assert {item["explaining_permutation"] for item in payload["definitional_items"]} == {
        "P-OCC-DENOM-ROOMS"
    }
    assert all(item["proposed_correction"] is None for item in payload["definitional_items"]), (
        "proposing a correction for a definitional variance tells a hotel to change a number that "
        "is not wrong"
    )


def test_the_permutation_id_exists_in_policy(permutations: PermutationView) -> None:
    """A fixture naming a permutation policy does not declare would expect an unreachable cause."""
    for spec in FIXTURES:
        if spec.mutation.permutation is not None:
            assert permutations.require(spec.mutation.permutation).applies_to

    with pytest.raises(FixtureError, match="declares no permutation"):
        permutations.require("P-INVENTED")


# ── the document and the expectation agree ───────────────────────────────────


def test_the_workbook_carries_the_figure_the_expectation_claims(
    built: tuple[FixtureBuild, ...],
) -> None:
    """The seam that matters: the cell the expectation cites holds the value it says it does.

    Both come from the same claim table, which is what makes them unable to disagree. This reads
    the rendered file back, because "unable to disagree" is an argument about the code and this is
    the fact.
    """
    payload = _payload(built, "F2")
    build = next(b for b in built if b.fixture_id == "F2")

    sheets = load_workbook(build.path / "submission" / "claims_2026-Q1.xlsx", data_only=True)
    for finding in payload["findings"]:
        excel = finding["evidence"]["excel"]
        cell = sheets[excel["sheet"]][excel["cell"]].value
        assert Decimal(str(cell)) == Decimal(finding["claimed"])


def test_the_untargeted_figures_are_left_alone(
    built: tuple[FixtureBuild, ...], table: tuple[ClaimCell, ...]
) -> None:
    """A mutation that moved a figure it was not aimed at would plant errors nobody declared."""
    build = next(b for b in built if b.fixture_id == "F2")
    sheets = load_workbook(build.path / "submission" / "claims_2026-Q1.xlsx", data_only=True)

    untouched = [
        claim
        for claim in table
        if claim.value_key not in (None, "DE") or claim.metric == "room_nights_sold"
    ]
    assert len(untouched) > 50, "this would pass vacuously on an empty list"
    for claim in untouched:
        assert sheets[claim.sheet][claim.cell].value == claim.value


# ── the preconditions the roll-up rests on ───────────────────────────────────


def test_every_additive_quarter_claim_is_the_sum_of_its_months(
    table: tuple[ClaimCell, ...],
) -> None:
    """The roll-up rewrites a quarter claim as a sum, which is only right if it already is one.

    Occupancy is excluded, and has to be: a quarter's occupancy is a ratio over the quarter, not
    three percentages added up. Asserted rather than assumed, because the roll-up would otherwise
    plant a second, undeclared error in whichever metric stopped being additive.
    """
    checked = 0
    for claim in table:
        if claim.period != QUARTER or claim.metric.endswith("_pct"):
            continue
        months = sum(
            other.value
            for other in table
            if other.metric == claim.metric
            and other.value_key == claim.value_key
            and other.period in MONTHS
        )
        assert claim.value == months, f"{claim.key} is not the sum of its months"
        checked += 1
    assert checked > 20, "the additive quarter claims were not found, so nothing was checked"


def test_the_roll_up_moves_with_the_month(
    table: tuple[ClaimCell, ...], policy: PolicyView, permutations: PermutationView
) -> None:
    """The mutated workbook has to stay consistent with itself.

    Left alone, the quarter column would no longer equal the three months beside it, and
    `tda.excel.selfcheck` would report that before any comparison happened. The fixture would then
    be exercising the self-consistency check rather than reconciliation, and it would look fine.
    """
    after = by_key(mutated(fixture("F2"), table, policy, permutations))
    key = "guests_by_nationality:2026-Q1:nationality_iso2=DE"
    months = sum(
        after[f"guests_by_nationality:{month}:nationality_iso2=DE"].value for month in MONTHS
    )

    assert after[key].value == months
    assert after[key].value != by_key(table)[key].value


def test_a_quarter_only_mutation_is_not_undone_by_the_roll_up(
    table: tuple[ClaimCell, ...], policy: PolicyView, permutations: PermutationView
) -> None:
    """A hotel that mistypes only its roll-up is a different fixture, and must stay one."""
    spec = FixtureSpec(
        fixture_id="F2",
        why=fixture("F2").why,
        mutation=Mutation(
            kind=MutationKind.TRANSPOSE_DIGITS,
            target=Target("guests_by_nationality", QUARTER, "DE"),
        ),
    )
    after = by_key(mutated(spec, table, policy, permutations))
    key = f"guests_by_nationality:{QUARTER}:nationality_iso2=DE"

    assert after[key].value == transpose_digits(by_key(table)[key].value)


# ── the mutation primitives ──────────────────────────────────────────────────


def test_the_transposition_swaps_the_leading_digits() -> None:
    assert transpose_digits(83) == 38
    assert transpose_digits(219) == 129


@pytest.mark.parametrize(
    ("value", "reason"),
    [(7, "one digit"), (44, "unchanged"), (69.09, "not a whole number")],
)
def test_the_transposition_refuses_what_it_cannot_damage(value: float, reason: str) -> None:
    """Each of these would produce a fixture whose planted error is not an error."""
    with pytest.raises(FixtureError, match=reason):
        transpose_digits(value)


def test_a_control_fixture_may_not_name_a_target() -> None:
    with pytest.raises(FixtureError, match="mutates nothing"):
        Mutation(kind=MutationKind.NONE, target=Target("occupancy_pct"))


def test_a_recomputation_without_a_permutation_is_refused() -> None:
    with pytest.raises(FixtureError, match="disagree"):
        Mutation(kind=MutationKind.RECOMPUTE_UNDER, target=Target("occupancy_pct"))


def test_a_permutation_on_anything_but_a_recomputation_is_refused() -> None:
    with pytest.raises(FixtureError, match="disagree"):
        Mutation(
            kind=MutationKind.TRANSPOSE_DIGITS,
            target=Target("guests_by_nationality", "2026-01", "DE"),
            permutation="P-OCC-DENOM-ROOMS",
        )


def test_a_dimension_mutation_must_name_a_dimension_value() -> None:
    with pytest.raises(FixtureError, match="must name one"):
        Mutation(kind=MutationKind.DELETE_DIMENSION, target=Target("guests_by_nationality"))


def test_the_deletion_arm_removes_every_claim_for_the_country(
    table: tuple[ClaimCell, ...], policy: PolicyView, permutations: PermutationView
) -> None:
    """F4's arm, ahead of F4. The set difference then does the classifying on its own."""
    spec = FixtureSpec(
        fixture_id="F4",
        why="A deleted nationality row is the completeness case, and it needs no value to compare.",
        mutation=Mutation(
            kind=MutationKind.DELETE_DIMENSION,
            target=Target("guests_by_nationality", value_key="IS"),
        ),
    )
    after = mutated(spec, table, policy, permutations)

    assert any(claim.value_key == "IS" for claim in table)
    assert not any(claim.value_key == "IS" for claim in after)
    assert len(after) == len(table) - 2, "Iceland appears in March and in the quarter total"


def test_a_deleted_row_derives_completeness_findings_with_no_cell(
    table: tuple[ClaimCell, ...],
    truth: dict[str, float],
    policy: PolicyView,
    permutations: PermutationView,
    view: ClassificationView,
) -> None:
    """The set difference classifies a structural absence without consulting the mutation at all.

    Which is the property worth pinning: `derive` reaches the mutation only for a value mismatch,
    so a kind with no entry in `CLASS_BY_KIND` still derives a complete expectation here.
    """
    spec = FixtureSpec(
        fixture_id="F4",
        why="A deleted nationality row is the completeness case, and it needs no value to compare.",
        mutation=Mutation(
            kind=MutationKind.DELETE_DIMENSION,
            target=Target("guests_by_nationality", value_key="IS"),
        ),
    )
    payload = _derive(spec, table, truth, policy, permutations, view)

    assert {f["key"] for f in payload["findings"]} == {
        "guests_by_nationality:2026-03:nationality_iso2=IS",
        f"guests_by_nationality:{QUARTER}:nationality_iso2=IS",
    }
    assert {f["variance_class"] for f in payload["findings"]} == {"V5"}
    assert all(f["claimed"] is None for f in payload["findings"])
    assert all(f["evidence"]["excel"] == "not_reached" for f in payload["findings"])
    assert all(f["difference"] is None for f in payload["findings"])


def test_a_claim_nothing_supports_derives_an_orphan_with_no_source(
    table: tuple[ClaimCell, ...],
    truth: dict[str, float],
    policy: PolicyView,
    view: ClassificationView,
) -> None:
    """D-MAT-05, the other direction of the set difference, and the more serious one.

    Iceland travelled in March only, so a January figure for Iceland is a claim the reports support
    no value for at all. Derived here against the **control** spec, whose kind has no entry in
    `CLASS_BY_KIND`: the structural arms have to classify this without the mutation being consulted,
    which is the property the whole derivation rests on.
    """
    march = next(claim for claim in table if claim.value_key == "IS" and claim.period == "2026-03")
    invented = dataclasses.replace(march, period="2026-01", cell="B11")

    payload: dict[str, Any] = expectation(fixture("F1"), (*table, invented), truth, policy, view)

    assert {f["key"] for f in payload["findings"]} == {
        "guests_by_nationality:2026-01:nationality_iso2=IS"
    }
    finding = payload["findings"][0]
    assert finding["variance_class"] == "V5"
    assert finding["computed"] is None
    assert finding["claimed"] == "3"
    assert finding["evidence"] == {
        "excel": {"sheet": "Nationality", "cell": "B11"},
        "source": "not_reached",
    }


def test_a_relabelling_changes_the_printed_label_and_nothing_else(
    table: tuple[ClaimCell, ...], policy: PolicyView, permutations: PermutationView
) -> None:
    """The mutation engine's half of F5 and F6: the label moves, the value and value_key do not.

    Resolvability is not this function's decision, so one spec exercises both directions of it by
    naming a label that happens not to matter here: what is asserted is that the claim's value
    survived the relabelling untouched, which is the property that lets `resolvable=True` derive an
    empty expectation later without a second line of mutation code.
    """
    spec = FixtureSpec(
        fixture_id="F5",
        why="A label the workbook prints differently for the same country.",
        mutation=Mutation(
            kind=MutationKind.RELABEL,
            target=Target("guests_by_nationality", value_key="CZ"),
            new_label="Czechia",
            resolvable=True,
        ),
    )
    before = {cell.key: cell.value for cell in table if cell.value_key == "CZ"}
    after = mutated(spec, table, policy, permutations)

    relabelled_cells = [cell for cell in after if cell.value_key == "CZ"]
    assert relabelled_cells, "the mutation must not remove the claims it relabels"
    assert all(cell.label == "Czechia" for cell in relabelled_cells)
    assert {cell.key: cell.value for cell in relabelled_cells} == before
    # Every other country's claims are untouched, not merely unequal in count.
    untouched = [cell for cell in after if cell.value_key != "CZ"]
    assert untouched == [cell for cell in table if cell.value_key != "CZ"]


def test_deleting_a_dimension_removes_every_period_it_was_claimed_in(
    table: tuple[ClaimCell, ...], policy: PolicyView, permutations: PermutationView
) -> None:
    """F4's half of the mutation engine: the row is gone, not zeroed, and gone everywhere it was."""
    spec = FixtureSpec(
        fixture_id="F4",
        why="A country dropped from the workbook entirely.",
        mutation=Mutation(
            kind=MutationKind.DELETE_DIMENSION,
            target=Target("guests_by_nationality", value_key="KZ"),
        ),
    )
    before = [cell for cell in table if cell.value_key == "KZ"]
    assert len(before) > 1, "the fixture is only interesting if the country spans more than one row"

    after = mutated(spec, table, policy, permutations)
    assert not any(cell.value_key == "KZ" for cell in after)
    assert len(after) == len(table) - len(before)


def test_an_unreproducible_permutation_is_refused(
    table: tuple[ClaimCell, ...], policy: PolicyView, permutations: PermutationView
) -> None:
    """A numerator-side permutation needs the ledger, and this builder says so rather than guessing.

    The quiet alternative is planting a figure the named permutation does not reproduce, which
    derives a definitional finding whose explanation does not explain it.
    """
    spec = FixtureSpec(
        fixture_id="F3",
        why=fixture("F3").why,
        mutation=Mutation(
            kind=MutationKind.RECOMPUTE_UNDER,
            target=Target("occupancy_pct"),
            permutation="P-MONTH-ARRIVAL",
        ),
    )
    with pytest.raises(FixtureError, match="cannot reproduce"):
        mutated(spec, table, policy, permutations)


# ── the policy reads ─────────────────────────────────────────────────────────


def test_a_missing_classification_rule_is_an_error_naming_the_path(
    document: dict[str, object],
) -> None:
    amended = copy.deepcopy(document)
    del amended["materiality"]

    with pytest.raises(FixtureError, match=r"materiality\.severity_by_class"):
        load_classification_view(amended)


def test_the_derivation_reads_the_configured_tolerance(
    table: tuple[ClaimCell, ...],
    truth: dict[str, float],
    policy: PolicyView,
    permutations: PermutationView,
    document: dict[str, object],
) -> None:
    """Widen the percentage band past the planted difference and the findings must disappear.

    A derivation that compared for inequality rather than against the configured band would still
    report four definitional items here, which is the D-TOL-04 rule it would be breaking: anything
    inside tolerance is logged and never raised.
    """
    amended = copy.deepcopy(document)
    amended["tolerances"]["by_metric_type"]["percentage_points"]["value"] = 1e9  # type: ignore[index]

    payload = _derive(
        fixture("F3"), table, truth, policy, permutations, load_classification_view(amended)
    )
    assert payload["definitional_items"] == []
    assert payload["status"] == "PASS"


def test_an_unknown_tolerance_type_is_refused(document: dict[str, object]) -> None:
    amended = copy.deepcopy(document)
    amended["tolerances"]["by_metric_type"]["count"]["type"] = "relative"  # type: ignore[index]

    with pytest.raises(FixtureError, match="relative tolerance"):
        load_classification_view(amended)


# ── the citation, and the halt that has no fixture yet ───────────────────────


def test_the_inventory_sourced_metric_cites_the_inventory_reference() -> None:
    """D-RNA-01. Rooms available is a property attribute and is on no page.

    `room_nights_available` is the one in-scope metric absent from the reservation export, so a
    finding on it cites the CSV. Describing it as a PDF page would send a reviewer to a document
    the denominator is not in; `not_reached` would claim nothing was consulted when the inventory
    reference was.
    """
    assert _source_of("room_nights_available") == "inventory_ref"
    assert _source_of("guests_by_nationality") == "pdf_page"


def test_occupancy_cites_the_reservation_rows_it_was_computed_from() -> None:
    """Occupancy carries both citations and the reservation rows come first.

    Pinned here because it is the one that reads wrongly at a glance: the denominator is not in the
    PDF, so `inventory_ref` looks like the careful answer. What the engine attaches is
    `source_rows[0]`, and `_occupancy` appends the inventory days after the reservation rows.
    """
    assert _source_of("occupancy_pct") == "pdf_page"


def test_a_blocking_finding_stops_the_expectation_at_itself(view: ClassificationView) -> None:
    """The short-circuit, ahead of the fixture that needs it.

    A run halts as soon as a blocking finding exists, so an expectation listing one *and* the
    reconciliation findings that would have followed describes a run that both halted and finished.
    No correct pipeline can produce that, and the fixture would fail as though the system were
    broken.
    """
    findings: list[dict[str, object]] = [
        {"key": "a", "variance_class": "V1", "severity": "material"},
        {"key": "b", "variance_class": "V7", "severity": "blocking"},
        {"key": "c", "variance_class": "V1", "severity": "material"},
    ]
    kept = _halt_at_first_blocking(findings)

    assert [f["key"] for f in kept] == ["b"]
    assert _halt_at_first_blocking(findings[:1]) == findings[:1]
    assert _status(kept, []) == "HALTED"
    assert view.severity_of("V7") == "blocking", "policy still says V7 outranks everything"


def test_a_resolvable_relabelling_derives_nothing(
    table: tuple[ClaimCell, ...],
    truth: dict[str, float],
    policy: PolicyView,
    permutations: PermutationView,
    view: ClassificationView,
) -> None:
    """F5's whole claim: a label the lookup resolves must derive an empty expectation.

    Through `mutated()` and `expectation()` together, the same two functions a real fixture build
    calls, rather than asserted against a hand-built table: this is the property F5 exists to prove,
    not a property of the test's own setup.
    """
    spec = FixtureSpec(
        fixture_id="F5",
        why="A label the workbook prints differently for the same country.",
        mutation=Mutation(
            kind=MutationKind.RELABEL,
            target=Target("guests_by_nationality", value_key="CZ"),
            new_label="Czechia",
            resolvable=True,
        ),
    )
    relabelled_table = mutated(spec, table, policy, permutations)
    payload = expectation(spec, relabelled_table, truth, policy, view)

    assert payload["status"] == "PASS"
    assert payload["findings"] == []
    assert payload["definitional_items"] == []


def test_an_unresolvable_relabelling_derives_one_blocking_finding(
    table: tuple[ClaimCell, ...],
    truth: dict[str, float],
    policy: PolicyView,
    permutations: PermutationView,
    view: ClassificationView,
) -> None:
    """F6's whole claim, and the one expectation this package cannot reach by set difference alone.

    The relabelled table still carries PL's correct value under the correct key, so nothing about
    the arithmetic can tell this case apart from a clean one. `expectation` has to be told directly,
    via `mutation.resolvable`, which is exactly why that field exists on the spec rather than being
    inferred here.
    """
    spec = FixtureSpec(
        fixture_id="F6",
        why="A country label the committed lookup has never seen.",
        mutation=Mutation(
            kind=MutationKind.RELABEL,
            target=Target("guests_by_nationality", value_key="PL"),
            new_label="Freedonia",
            resolvable=False,
        ),
    )
    relabelled_table = mutated(spec, table, policy, permutations)
    payload = expectation(spec, relabelled_table, truth, policy, view)

    assert payload["status"] == "HALTED"
    assert payload["definitional_items"] == []
    findings = payload["findings"]
    assert isinstance(findings, list)
    assert len(findings) == 1
    finding = findings[0]
    assert isinstance(finding, dict)
    assert finding["variance_class"] == "V7"
    assert finding["severity"] == "blocking"
    assert finding["escalates_to"] == "human_review"
    evidence = finding["evidence"]
    assert isinstance(evidence, dict)
    assert evidence["source"] == "not_reached"
    excel = evidence["excel"]
    assert isinstance(excel, dict) and excel["sheet"] == "Nationality"


def test_the_status_rules_are_the_ones_the_graph_applies() -> None:
    """Most serious first, and PASS means zero of both rather than zero worth mentioning."""
    material: list[dict[str, object]] = [{"severity": "material"}]

    assert _status([], []) == "PASS"
    assert _status(material, []) == "FAIL"
    assert _status([], material) == "ESCALATED"
    assert _status(material, material) == "ESCALATED"
    blocking: list[dict[str, object]] = [{"severity": "blocking"}]
    assert _status(blocking, material) == "HALTED"


@pytest.mark.parametrize("fixture_id", [spec.fixture_id for spec in FIXTURES])
def test_a_fixture_submits_exactly_what_the_demo_hotel_submits(
    built: tuple[FixtureBuild, ...], fixture_id: str
) -> None:
    """A fixture is a whole submission, with one document changed and the rest byte-identical.

    The pipeline reads the three reports and the inventory reference as well as the workbook, so a
    fixture missing one of them is not a harder case, it is an unusable one. And every file except
    the workbook has to match the demo byte for byte: a re-rendered PDF would make the fixture
    sensitive to a reportlab upgrade, and six fixtures would then change at once for a reason that
    has nothing to do with any mutation.
    """
    build = next(b for b in built if b.fixture_id == fixture_id)
    source = DEMO / "submission"
    produced = build.path / "submission"

    assert {p.name for p in produced.iterdir()} == {p.name for p in source.iterdir()}
    for path in sorted(source.iterdir()):
        if path.name == "claims_2026-Q1.xlsx":
            continue
        assert (produced / path.name).read_bytes() == path.read_bytes()


def test_one_fixture_directory_per_spec(built: tuple[FixtureBuild, ...]) -> None:
    assert [build.fixture_id for build in built] == [spec.fixture_id for spec in FIXTURES]
    for build in built:
        assert (build.path / "README.md").is_file()
        assert (build.path / "manifest.json").read_bytes() == (DEMO / "manifest.json").read_bytes()
        assert not (build.path / "ground_truth").exists(), (
            "a fixture that carried the answers beside the questions is one wrong path away from a "
            "verification that had seen them"
        )


def test_an_unknown_fixture_id_names_the_set(
    truth: dict[str, float],
    policy: PolicyView,
    permutations: PermutationView,
    view: ClassificationView,
    tmp_path: Path,
) -> None:
    with pytest.raises(FixtureError, match="this set carries"):
        fixture("F9")

    # And a spec built for a fixture id the schema forbids fails on write rather than on read.
    spec = FixtureSpec(fixture_id="F9", why=fixture("F1").why, mutation=Mutation(MutationKind.NONE))
    with pytest.raises(FixtureError, match=r"expected\.schema\.json"):
        materialise_one(spec, DEMO, tmp_path, truth, policy, permutations, view)
